"""
Export Handlers for 2026 Invoice Extraction System

Supports:
- JSON with full validation metadata and confidence per field
- Excel with 5 sheets (Headers, Line Items, Validation, Confidence, Metadata)
- CSV (flat headers + separate line items)

Author: ML Engineering Team
Version: 2.0.0
"""

import json
import csv
import logging
import os
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional

from src.models.base_extractor import ExtractionOutput

logger = logging.getLogger("invoice_extraction.export")


class JSONExporter2026:
    """Export extraction results to JSON with full metadata."""

    def export(
        self,
        extractions: List[ExtractionOutput],
        output_path: str,
        validation_reports: Optional[List[Dict[str, Any]]] = None,
        metadata: Optional[Dict[str, Any]] = None,
    ) -> str:
        """
        Export extractions to a structured JSON file.

        Returns the output file path.
        """
        os.makedirs(os.path.dirname(output_path), exist_ok=True)

        output = {
            "metadata": {
                "export_timestamp": datetime.now().isoformat(),
                "system_version": "2.0.0",
                "total_invoices": len(extractions),
                **(metadata or {}),
            },
            "invoices": [],
        }

        for i, extraction in enumerate(extractions):
            invoice_data = {
                "index": i + 1,
                "source_file": extraction.source_file,
                "page_number": extraction.page_number,
                "headers": {},
                "line_items": [],
                "quality_metadata": extraction.quality_metadata,
                "overall_confidence": extraction.calculate_overall_confidence(),
                "extraction_model": extraction.model_used,
                "processing_time_ms": extraction.processing_time_ms,
            }

            # Header fields with confidence
            for field_name in ExtractionOutput.HEADER_FIELDS:
                fv = extraction.get_field(field_name)
                if fv:
                    invoice_data["headers"][field_name] = {
                        "value": fv.value,
                        "confidence": round(fv.confidence, 1),
                        "uncertain": fv.uncertain,
                        "verified": fv.verified,
                        "validation_method": fv.validation_method,
                        "source_model": fv.source_model,
                    }

            # Line items
            for item in extraction.line_items:
                item_data = {
                    "line_number": item.line_number,
                    "description": {
                        "value": item.description.value,
                        "confidence": round(item.description.confidence, 1),
                    },
                    "quantity": {
                        "value": item.quantity.value,
                        "confidence": round(item.quantity.confidence, 1),
                    },
                    "unit_price": {
                        "value": item.unit_price.value,
                        "confidence": round(item.unit_price.confidence, 1),
                    },
                    "line_total": {
                        "value": item.line_total.value,
                        "confidence": round(item.line_total.confidence, 1),
                    },
                }
                if item.unit_of_measure.value:
                    item_data["unit_of_measure"] = {
                        "value": item.unit_of_measure.value,
                        "confidence": round(item.unit_of_measure.confidence, 1),
                    }
                if item.tax_rate.value:
                    item_data["tax_rate"] = {
                        "value": item.tax_rate.value,
                        "confidence": round(item.tax_rate.confidence, 1),
                    }
                if item.item_code.value:
                    item_data["item_code"] = {
                        "value": item.item_code.value,
                        "confidence": round(item.item_code.confidence, 1),
                    }
                invoice_data["line_items"].append(item_data)

            # Validation report
            if validation_reports and i < len(validation_reports):
                invoice_data["validation_report"] = validation_reports[i]
            elif extraction.validation_results:
                invoice_data["validation_report"] = extraction.validation_results

            output["invoices"].append(invoice_data)

        # Summary stats
        output["summary"] = self._compute_summary(extractions)

        with open(output_path, "w", encoding="utf-8") as f:
            json.dump(output, f, indent=2, ensure_ascii=False, default=str)

        logger.info(f"JSON export: {len(extractions)} invoices → {output_path}")
        return output_path

    def _compute_summary(self, extractions: List[ExtractionOutput]) -> Dict:
        """Compute summary statistics."""
        confidences = [e.calculate_overall_confidence() for e in extractions]
        total_items = sum(len(e.line_items) for e in extractions)
        uncertain_fields = 0
        verified_fields = 0

        for e in extractions:
            for field_name in ExtractionOutput.HEADER_FIELDS:
                fv = e.get_field(field_name)
                if fv and fv.value is not None:
                    if fv.uncertain:
                        uncertain_fields += 1
                    if fv.verified:
                        verified_fields += 1

        return {
            "total_invoices": len(extractions),
            "total_line_items": total_items,
            "avg_confidence": round(
                sum(confidences) / len(confidences), 1
            ) if confidences else 0,
            "min_confidence": round(min(confidences), 1) if confidences else 0,
            "max_confidence": round(max(confidences), 1) if confidences else 0,
            "uncertain_fields": uncertain_fields,
            "verified_fields": verified_fields,
        }


class ExcelExporter2026:
    """
    Export to Excel with 5 sheets:
    1. Headers - All header fields with values
    2. Line Items - All line items across invoices
    3. Validation Report - Per-invoice validation results
    4. Confidence Scores - Field-level confidence breakdown
    5. Model Metadata - Model info, timing, quality assessment
    """

    def export(
        self,
        extractions: List[ExtractionOutput],
        output_path: str,
        validation_reports: Optional[List[Dict[str, Any]]] = None,
    ) -> str:
        """Export to multi-sheet Excel file."""
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        except ImportError:
            logger.error("openpyxl not installed. Run: pip install openpyxl")
            raise

        os.makedirs(os.path.dirname(output_path), exist_ok=True)

        wb = openpyxl.Workbook()

        # Sheet 1: Headers
        self._create_headers_sheet(wb, extractions)

        # Sheet 2: Line Items
        self._create_line_items_sheet(wb, extractions)

        # Sheet 3: Validation Report
        self._create_validation_sheet(wb, extractions, validation_reports)

        # Sheet 4: Confidence Scores
        self._create_confidence_sheet(wb, extractions)

        # Sheet 5: Metadata
        self._create_metadata_sheet(wb, extractions)

        # Remove default sheet if others exist
        if "Sheet" in wb.sheetnames and len(wb.sheetnames) > 1:
            del wb["Sheet"]

        wb.save(output_path)
        logger.info(f"Excel export: {len(extractions)} invoices → {output_path}")
        return output_path

    def _create_headers_sheet(self, wb, extractions):
        """Sheet 1: Header fields."""
        ws = wb.create_sheet("Headers")

        # Header row
        columns = [
            "Invoice #", "Source File", "Invoice Number", "Invoice Date",
            "Due Date", "Vendor Name", "Vendor Address", "Vendor Email",
            "Vendor Phone", "Customer Name", "Customer Address",
            "Currency", "Subtotal", "Tax Amount", "Shipping", "Total Amount",
            "Overall Confidence (%)", "Model Used",
        ]
        self._write_header_row(ws, columns)

        for i, extraction in enumerate(extractions):
            row = [
                i + 1,
                extraction.source_file or "",
                self._val(extraction.invoice_number),
                self._val(extraction.invoice_date),
                self._val(extraction.due_date),
                self._val(extraction.vendor_name),
                self._val(extraction.vendor_address),
                self._val(extraction.vendor_email),
                self._val(extraction.vendor_phone),
                self._val(extraction.customer_name),
                self._val(extraction.customer_address),
                self._val(extraction.currency),
                self._val(extraction.subtotal),
                self._val(extraction.tax_amount),
                self._val(extraction.shipping),
                self._val(extraction.total_amount),
                round(extraction.calculate_overall_confidence(), 1),
                extraction.model_used or "",
            ]
            ws.append(row)

        self._auto_width(ws)

    def _create_line_items_sheet(self, wb, extractions):
        """Sheet 2: All line items."""
        ws = wb.create_sheet("Line Items")

        columns = [
            "Invoice #", "Source File", "Line #", "Description",
            "Quantity", "Unit Price", "Line Total",
            "Item Code", "Unit of Measure", "Tax Rate",
            "Desc Confidence (%)", "Qty Confidence (%)",
            "Price Confidence (%)", "Total Confidence (%)",
        ]
        self._write_header_row(ws, columns)

        for i, extraction in enumerate(extractions):
            for item in extraction.line_items:
                row = [
                    i + 1,
                    extraction.source_file or "",
                    item.line_number,
                    self._val(item.description),
                    self._val(item.quantity),
                    self._val(item.unit_price),
                    self._val(item.line_total),
                    self._val(item.item_code),
                    self._val(item.unit_of_measure),
                    self._val(item.tax_rate),
                    round(item.description.confidence, 1),
                    round(item.quantity.confidence, 1),
                    round(item.unit_price.confidence, 1),
                    round(item.line_total.confidence, 1),
                ]
                ws.append(row)

        self._auto_width(ws)

    def _create_validation_sheet(self, wb, extractions, validation_reports):
        """Sheet 3: Validation results."""
        ws = wb.create_sheet("Validation Report")

        columns = [
            "Invoice #", "Source File", "Overall Passed", "Overall Score",
            "Decision", "RAG Score", "Multi-Agent Score",
            "QAE Score", "Neurosymbolic Score", "Confidence Score",
            "Cross-Model Score", "Issues Count", "Corrections Count",
        ]
        self._write_header_row(ws, columns)

        for i, extraction in enumerate(extractions):
            vr = None
            if validation_reports and i < len(validation_reports):
                vr = validation_reports[i]
            elif extraction.validation_results:
                vr = extraction.validation_results

            if vr:
                scores = vr.get("layer_scores", {})
                row = [
                    i + 1,
                    extraction.source_file or "",
                    str(vr.get("overall_passed", "")),
                    round(vr.get("overall_score", 0), 3),
                    vr.get("overall_decision", ""),
                    round(scores.get("rag", 0), 3),
                    round(scores.get("multi_agent", 0), 3),
                    round(scores.get("qae", 0), 3),
                    round(scores.get("neurosymbolic", 0), 3),
                    round(scores.get("confidence", 0), 3),
                    round(scores.get("cross_model", 0), 3),
                    len(vr.get("all_issues", [])),
                    len(vr.get("corrections_applied", [])),
                ]
            else:
                row = [i + 1, extraction.source_file or ""] + ["N/A"] * 11

            ws.append(row)

        self._auto_width(ws)

    def _create_confidence_sheet(self, wb, extractions):
        """Sheet 4: Per-field confidence breakdown."""
        ws = wb.create_sheet("Confidence Scores")

        columns = ["Invoice #", "Source File"]
        for field_name in ExtractionOutput.HEADER_FIELDS:
            columns.append(f"{field_name} (%)")
            columns.append(f"{field_name} uncertain")
            columns.append(f"{field_name} verified")
        self._write_header_row(ws, columns)

        for i, extraction in enumerate(extractions):
            row = [i + 1, extraction.source_file or ""]
            for field_name in ExtractionOutput.HEADER_FIELDS:
                fv = extraction.get_field(field_name)
                if fv and fv.value is not None:
                    row.extend([
                        round(fv.confidence, 1),
                        str(fv.uncertain),
                        str(fv.verified),
                    ])
                else:
                    row.extend(["N/A", "N/A", "N/A"])
            ws.append(row)

        self._auto_width(ws)

    def _create_metadata_sheet(self, wb, extractions):
        """Sheet 5: Model and processing metadata."""
        ws = wb.create_sheet("Model Metadata")

        columns = [
            "Invoice #", "Source File", "Model Used", "Page Number",
            "Processing Time (ms)", "Image Quality",
            "Quality Score", "Export Timestamp",
        ]
        self._write_header_row(ws, columns)

        for i, extraction in enumerate(extractions):
            qm = extraction.quality_metadata or {}
            row = [
                i + 1,
                extraction.source_file or "",
                extraction.model_used or "",
                extraction.page_number or 1,
                extraction.processing_time_ms or 0,
                qm.get("overall_quality", "N/A"),
                round(qm.get("quality_score", 0), 3)
                if qm.get("quality_score") else "N/A",
                datetime.now().isoformat(),
            ]
            ws.append(row)

        self._auto_width(ws)

    def _write_header_row(self, ws, columns):
        """Write styled header row."""
        try:
            from openpyxl.styles import Font, PatternFill, Alignment

            header_font = Font(bold=True, size=11, color="FFFFFF")
            header_fill = PatternFill(
                start_color="2F5496", end_color="2F5496", fill_type="solid"
            )
            for col_idx, col_name in enumerate(columns, 1):
                cell = ws.cell(row=1, column=col_idx, value=col_name)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", wrap_text=True)
        except ImportError:
            ws.append(columns)

    def _auto_width(self, ws):
        """Auto-adjust column widths."""
        for column in ws.columns:
            max_length = 0
            col_letter = None
            for cell in column:
                if col_letter is None:
                    col_letter = cell.column_letter
                try:
                    if cell.value:
                        max_length = max(max_length, len(str(cell.value)))
                except Exception:
                    pass
            if col_letter:
                ws.column_dimensions[col_letter].width = min(max_length + 3, 50)

    def _val(self, field_value):
        """Extract value from FieldValue."""
        if hasattr(field_value, "value"):
            return field_value.value if field_value.value is not None else ""
        return field_value or ""


class CSVExporter2026:
    """Export to CSV (flat format)."""

    def export_headers(
        self,
        extractions: List[ExtractionOutput],
        output_path: str,
    ) -> str:
        """Export header fields to CSV."""
        os.makedirs(os.path.dirname(output_path), exist_ok=True)

        columns = [
            "index", "source_file", "invoice_number", "invoice_date",
            "due_date", "vendor_name", "vendor_address", "vendor_email",
            "vendor_phone", "customer_name", "customer_address",
            "currency", "subtotal", "tax_amount", "shipping", "total_amount",
            "overall_confidence", "model_used",
        ]

        with open(output_path, "w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=columns)
            writer.writeheader()

            for i, extraction in enumerate(extractions):
                row = {
                    "index": i + 1,
                    "source_file": extraction.source_file or "",
                    "overall_confidence": round(
                        extraction.calculate_overall_confidence(), 1
                    ),
                    "model_used": extraction.model_used or "",
                }
                for field_name in ExtractionOutput.HEADER_FIELDS:
                    fv = extraction.get_field(field_name)
                    row[field_name] = fv.value if fv and fv.value is not None else ""
                writer.writerow(row)

        logger.info(f"CSV headers export: {len(extractions)} → {output_path}")
        return output_path

    def export_line_items(
        self,
        extractions: List[ExtractionOutput],
        output_path: str,
    ) -> str:
        """Export line items to CSV."""
        os.makedirs(os.path.dirname(output_path), exist_ok=True)

        columns = [
            "invoice_index", "source_file", "line_number",
            "description", "quantity", "unit_price", "line_total",
            "item_code", "unit_of_measure", "tax_rate",
        ]

        with open(output_path, "w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=columns)
            writer.writeheader()

            for i, extraction in enumerate(extractions):
                for item in extraction.line_items:
                    writer.writerow({
                        "invoice_index": i + 1,
                        "source_file": extraction.source_file or "",
                        "line_number": item.line_number,
                        "description": item.description.value or "",
                        "quantity": item.quantity.value or "",
                        "unit_price": item.unit_price.value or "",
                        "line_total": item.line_total.value or "",
                        "item_code": item.item_code.value or "",
                        "unit_of_measure": item.unit_of_measure.value or "",
                        "tax_rate": item.tax_rate.value or "",
                    })

        logger.info(f"CSV line items export → {output_path}")
        return output_path


class ExportManager2026:
    """Manages all export formats."""

    def __init__(self, config: Dict[str, Any]):
        self.config = config
        output_config = config.get("output", {})
        self.formats = output_config.get("formats", ["json", "excel"])
        self.output_dir = output_config.get(
            "output_dir", "outputs/extractions"
        )

        self.json_exporter = JSONExporter2026()
        self.excel_exporter = ExcelExporter2026()
        self.csv_exporter = CSVExporter2026()

    def export_all(
        self,
        extractions: List[ExtractionOutput],
        validation_reports: Optional[List[Dict[str, Any]]] = None,
        metadata: Optional[Dict[str, Any]] = None,
        timestamp: Optional[str] = None,
    ) -> Dict[str, str]:
        """
        Export to all configured formats.

        Returns dict of format → output_path.
        """
        if not timestamp:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")

        results = {}

        if "json" in self.formats:
            path = os.path.join(
                self.output_dir, f"invoice_extractions_{timestamp}.json"
            )
            results["json"] = self.json_exporter.export(
                extractions, path, validation_reports, metadata
            )

        if "excel" in self.formats:
            path = os.path.join(
                self.output_dir, f"invoice_extractions_{timestamp}.xlsx"
            )
            results["excel"] = self.excel_exporter.export(
                extractions, path, validation_reports
            )

        if "csv" in self.formats:
            headers_path = os.path.join(
                self.output_dir, f"invoice_headers_{timestamp}.csv"
            )
            items_path = os.path.join(
                self.output_dir, f"invoice_line_items_{timestamp}.csv"
            )
            results["csv_headers"] = self.csv_exporter.export_headers(
                extractions, headers_path
            )
            results["csv_line_items"] = self.csv_exporter.export_line_items(
                extractions, items_path
            )

        logger.info(f"Exported to {len(results)} files: {list(results.keys())}")
        return results
